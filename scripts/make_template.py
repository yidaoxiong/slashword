#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成「词库导入模板.xlsx」，放到 public/ 下，家长页可以直接点链接下载。

模板的表头就是导入时识别的**标准列名**，所以改这里的表头 = 改导入规则，
两边要一起改（前端别名表在 src/lib/importer.ts 的 COLUMNS）。
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "public" / "词库导入模板.xlsx"

HEADERS = [
    ("单词", True),
    ("音标", False),
    ("词性", False),
    ("中文含义", True),
    ("英文例句", False),
    ("例句中文", False),
    ("单元", False),
    ("单元主题", False),
    ("课文", False),
    ("词汇类别", False),
]

# 示例：一个「食物」主题的六词小词库，覆盖常见填写情况
SAMPLE = [
    ("sandwich", "英 /ˈsænwɪdʒ/ 美 /ˈsænwɪdʒ/", "n.", "三明治",
     "I had a sandwich for lunch.", "我午饭吃了个三明治。", "Unit 1", "食物", "Lesson 1", "Target"),
    ("soup", "/suːp/", "n.", "汤",
     "This soup is too hot.", "这汤太烫了。", "Unit 1", "食物", "Lesson 1", "Target"),
    ("salad", "英 /ˈsæləd/ 美 /ˈsæləd/", "n.", "沙拉",
     "She made a green salad.", "她做了一份蔬菜沙拉。", "Unit 1", "食物", "Lesson 1", "Target"),
    ("cookie", "/ˈkʊki/", "n.", "曲奇饼干",
     "Would you like a cookie?", "你想来块曲奇吗？", "Unit 1", "食物", "Lesson 2", "Target"),
    ("orange juice", "/ˈɒrɪndʒ dʒuːs/", "n.", "橙汁",
     "I drink orange juice every morning.", "我每天早上喝橙汁。", "Unit 2", "饮料", "Lesson 3", "Context"),
    ("Would you like ... ?", "/wʊd juː laɪk/", "句型", "你想要……吗？",
     "Would you like some tea?", "你想喝点茶吗？", "Unit 2", "饮料", "Lesson 3", "Extension"),
]

GUIDE = [
    ("怎么填", ""),
    ("1. 必填两列", "「单词」和「中文含义」。缺单词的行会被自动跳过。"),
    ("2. 音标", "可以只写一个 /suːp/，也可以写「英 /a/ 美 /b/」两种都给。留空也能用，只是卡片上不显示。"),
    ("3. 例句", "有例句就有「例句训练」这一关的材料，没有也能学，只是那一关会退化成看中文选英文。"),
    ("4. 单元", "填 Unit 1 / 第一单元 / 1 都行，系统会按顺序排。留空就当成一个整体，学的时候不分单元。"),
    ("5. 词汇类别", "比如厚海词表分 Target / Context / Extension。没有就整列留空。"),
    ("", ""),
    ("注意", ""),
    ("括号注释", "单词里别写 maths (=mathematics) 这种，写了也会被自动剥掉，注释单独存起来。"),
    ("二选一", "turn left/right 这种只取前一半作为拼写答案。"),
    ("重复词", "同一个词出现多次，后面的会被跳过（记忆卡按词归一，重复没意义）。"),
    ("", ""),
    ("导入之后", ""),
    ("发音", "新词库导入后立刻能学，发音先用设备自带的语音。"),
    ("", "想要和内置词库一样的高品质发音，在家长页点「导出词库文件」，"),
    ("", "把下载到的 json 放进 word-app/data/ 目录，再跑："),
    ("", "    npm run import-book data/词库id.json"),
    ("", "    npm run audio"),
    ("", "    npm run deploy"),
    ("", "之后这个词库就有了微软神经网络发音，并且不会再因为清缓存丢失。"),
]


def build():
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "单词表"

    head_fill = PatternFill("solid", fgColor="7132F5")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    req_font = Font(color="FFFFFF", bold=True, size=11)

    for i, (name, required) in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=name + (" *" if required else ""))
        c.fill = head_fill
        c.font = req_font if required else head_font
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    for r, row in enumerate(SAMPLE, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)

    widths = [20, 26, 8, 20, 38, 24, 12, 12, 12, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"

    g = wb.create_sheet("填写说明")
    g.column_dimensions["A"].width = 16
    g.column_dimensions["B"].width = 78
    for r, (a, b) in enumerate(GUIDE, start=1):
        ca = g.cell(row=r, column=1, value=a)
        cb = g.cell(row=r, column=2, value=b)
        if b == "" and a != "":
            ca.font = Font(bold=True, size=11, color="7132F5")
        cb.alignment = Alignment(wrap_text=True, vertical="top")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"-> 已生成 {OUT}（{OUT.stat().st_size / 1024:.1f} KB）")


if __name__ == "__main__":
    build()
